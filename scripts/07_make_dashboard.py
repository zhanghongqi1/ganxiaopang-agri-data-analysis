# -*- coding: utf-8 -*-
"""
07_make_dashboard.py — 生成多 sheet Excel 经营分析看板（openpyxl）

数据源：
  - logs/analysis_results.json        （05_run_queries.py 产出的核心指标）
  - data/processed/analysis_*.csv     （各分析模块明细表）
  - reports/figures/fig*.png          （06_make_figures.py 产出的图表，嵌入看板）

输出：
  - dashboard/甘小胖经营分析看板.xlsx （6 个 sheet：核心指标/品类ROI/用户分层/直播漏斗/供应链/门店测算）

声明：本项目全部数据为 Faker 模拟数据，仅用于方法复现，不代表真实经营结果。
"""
import json
import struct
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
JSON_PATH = ROOT / "logs" / "analysis_results.json"
PROC_DIR = ROOT / "data" / "processed"
FIG_DIR = ROOT / "reports" / "figures"
OUT_DIR = ROOT / "dashboard"
OUT_DIR.mkdir(exist_ok=True)
OUT_PATH = OUT_DIR / "甘小胖经营分析看板.xlsx"

DISCLAIMER = "声明：本看板全部数据为 Faker 模拟数据，仅用于方法复现，不代表真实经营结果。"

# ---------------------------------------------------------------- 样式常量
C_PRIMARY = "2E7D32"      # 农业主题深绿
C_PRIMARY_DARK = "1B5E20"
C_ACCENT = "F9A825"       # 麦浪黄
C_HEADER_BG = PatternFill("solid", fgColor=C_PRIMARY)
C_TITLE_BG = PatternFill("solid", fgColor=C_PRIMARY_DARK)
C_ALT_BG = PatternFill("solid", fgColor="F1F8E9")     # 交替行浅绿
C_KPI_BG = PatternFill("solid", fgColor="E8F5E9")
C_WARN_BG = PatternFill("solid", fgColor="FDECEA")    # 亏损/预警浅红
C_GOOD_BG = PatternFill("solid", fgColor="E6F4EA")    # 达标浅绿
C_GRAY_BG = PatternFill("solid", fgColor="F5F5F5")

F_TITLE = Font(name="微软雅黑", size=16, bold=True, color="FFFFFF")
F_H1 = Font(name="微软雅黑", size=13, bold=True, color=C_PRIMARY_DARK)
F_HEADER = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
F_BODY = Font(name="微软雅黑", size=10)
F_BOLD = Font(name="微软雅黑", size=10, bold=True)
F_KPI_NUM = Font(name="微软雅黑", size=18, bold=True, color=C_PRIMARY_DARK)
F_KPI_LABEL = Font(name="微软雅黑", size=10, color="555555")
F_WARN = Font(name="微软雅黑", size=10, bold=True, color="C62828")
F_GOOD = Font(name="微软雅黑", size=10, bold=True, color="2E7D32")
F_DISCLAIMER = Font(name="微软雅黑", size=10, bold=True, color="C62828")
F_LINK = Font(name="微软雅黑", size=10, color="1565C0", underline="single")

THIN = Side(style="thin", color="BDBDBD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


# ---------------------------------------------------------------- 工具函数
def png_size(path: Path):
    """不依赖 PIL，直接解析 PNG 文件头 IHDR 获取像素尺寸。"""
    with open(path, "rb") as f:
        f.read(16)
        w, h = struct.unpack(">II", f.read(8))
    return w, h


def add_image(ws, fig_name: str, anchor_row: int, anchor_col: int = 1, max_w: int = 780):
    """按比例缩放嵌入 PNG 图表，返回图片占用的行数（用于下一张图定位）。"""
    path = FIG_DIR / fig_name
    if not path.exists():
        return 0
    w, h = png_size(path)
    scale = max_w / w
    img = XLImage(str(path))
    img.width = int(w * scale)
    img.height = int(h * scale)
    anchor = f"{get_column_letter(anchor_col)}{anchor_row}"
    ws.add_image(img, anchor)
    # 每行默认约 20px，预留行数
    return img.height // 20 + 2


# 列名 → 数字格式
def guess_fmt(col: str) -> str:
    if any(k in col for k in ["GMV", "金额", "毛利", "成本", "销售额", "净利", "运费",
                              "缺口", "摊销", "租金", "杂费", "投入", "采购额", "供货额",
                              "客单价", "UV价值", "坪效", "人效", "贡献金额"]):
        return "#,##0.00"
    if any(k in col for k in ["ROI", "时效", "频次", "单量", "单位运费", "平均金额"]):
        return "#,##0.00"
    if "_pct" in col or "率" in col or "占比" in col:
        return "0.00"
    if any(k in col for k in ["天数", "回收期", "日均单量"]):
        return "0.0"
    if any(k in col for k in ["数", "人次", "场次数", "SKU", "面积", "量_kg"]):
        return "#,##0"
    return "General"


def write_df(ws, df: pd.DataFrame, start_row: int, start_col: int = 1,
             warn_neg_col: str = None, roi_col: str = None) -> int:
    """把 DataFrame 写入 sheet（带表头样式/交替行/数字格式/条件着色），返回下一可用行。"""
    n_cols = len(df.columns)
    # 表头
    for j, col in enumerate(df.columns):
        c = ws.cell(row=start_row, column=start_col + j, value=str(col))
        c.font = F_HEADER
        c.fill = C_HEADER_BG
        c.alignment = CENTER
        c.border = BORDER
    # 数据行
    for i, (_, row) in enumerate(df.iterrows()):
        r = start_row + 1 + i
        for j, col in enumerate(df.columns):
            v = row[col]
            if pd.isna(v):
                v = "—"
            elif isinstance(v, float):
                v = round(v, 4)
            c = ws.cell(row=r, column=start_col + j, value=v)
            c.font = F_BODY
            c.border = BORDER
            c.alignment = RIGHT if isinstance(v, (int, float)) else LEFT
            if isinstance(v, (int, float)):
                c.number_format = guess_fmt(str(col))
            if i % 2 == 1:
                c.fill = C_ALT_BG
            # 条件着色
            if warn_neg_col and col == warn_neg_col and isinstance(v, (int, float)) and v < 0:
                c.font = F_WARN
                c.fill = C_WARN_BG
            if roi_col and col == roi_col and isinstance(v, (int, float)):
                c.font = F_GOOD if v >= 1 else F_WARN
                c.fill = C_GOOD_BG if v >= 1 else C_WARN_BG
    # 列宽
    for j, col in enumerate(df.columns):
        max_len = len(str(col))
        for i in range(min(len(df), 30)):
            v = df.iloc[i, j]
            s = f"{v:,.2f}" if isinstance(v, float) else str(v)
            max_len = max(max_len, len(s))
        width = min(max(max_len * 1.6 + 2, 10), 34)
        # 中文按双宽估算
        cn = sum(1 for ch in str(col) if ord(ch) > 127)
        width = min(max(width, (len(str(col)) + cn) * 1.4 + 4), 34)
        ws.column_dimensions[get_column_letter(start_col + j)].width = width
    return start_row + 1 + len(df)


def write_section_title(ws, row: int, text: str, n_cols: int = 8):
    c = ws.cell(row=row, column=1, value=text)
    c.font = F_H1
    c.alignment = LEFT
    return row + 1


def write_note(ws, row: int, text: str, font=None):
    c = ws.cell(row=row, column=1, value=text)
    c.font = font or Font(name="微软雅黑", size=9, color="777777")
    c.alignment = LEFT
    return row + 1


def load_csv(name: str) -> pd.DataFrame:
    return pd.read_csv(PROC_DIR / name, encoding="utf-8-sig")


def sheet_footer(ws, row: int):
    write_note(ws, row + 1, DISCLAIMER, F_DISCLAIMER)


# ---------------------------------------------------------------- 数据加载
with open(JSON_PATH, encoding="utf-8") as f:
    R = json.load(f)

df_cat = load_csv("analysis_a_category_roi.csv")
df_rfm = load_csv("analysis_b_user_rfm.csv")
df_funnel = load_csv("analysis_c_live_funnel.csv")
df_slot = load_csv("analysis_c_live_slot.csv")
df_mode = load_csv("analysis_d_delivery_mode.csv")
df_loss = load_csv("analysis_d_loss.csv")
df_store = load_csv("analysis_e_store_pnl.csv")
df_store_sum = load_csv("analysis_e_store_summary.csv")

wb = Workbook()

# ================================================================ Sheet1 核心指标
ws = wb.active
ws.title = "核心指标"
ws.sheet_view.showGridLines = False

ws.merge_cells("A1:J2")
c = ws["A1"]
c.value = "甘小胖 · 扎根富农产业链 经营数据分析看板"
c.font = F_TITLE
c.fill = C_TITLE_BG
c.alignment = CENTER
for row in ws["A1:J2"]:
    for cell in row:
        cell.fill = C_TITLE_BG

ws.merge_cells("A3:J3")
c = ws["A3"]
c.value = DISCLAIMER
c.font = F_DISCLAIMER
c.fill = C_WARN_BG
c.alignment = CENTER

# 目录
row = 5
row = write_section_title(ws, row, "📑 目录（点击跳转）")
toc = [
    ("① 品类结构与 ROI 分析", "品类ROI"),
    ("② 用户 RFM 分层与复购", "用户分层"),
    ("③ 直播转化漏斗与时段", "直播漏斗"),
    ("④ 供应链：配送模式与损耗", "供应链"),
    ("⑤ 门店盈利测算与盈亏平衡", "门店测算"),
]
for name, target in toc:
    c = ws.cell(row=row, column=1, value=name)
    c.hyperlink = f"#'{target}'!A1"
    c.font = F_LINK
    row += 1

# KPI 卡片区（2 行 × 4 列）
row += 1
row = write_section_title(ws, row, "📊 核心经营指标（2025 全年）")

cat = R["category"]
usr = R["user"]
sup = R["supply_chain"]
sto = R["store"]["summary"][0]
new_mode = [m for m in sup["mode_compare"] if m["配送模式"].startswith("新模式") and m["业务类型"] == "电商订单"][0]
old_mode = [m for m in sup["mode_compare"] if m["配送模式"].startswith("旧模式") and m["业务类型"] == "电商订单"][0]

kpis = [
    ("全渠道 GMV", f"¥{cat['total_gmv']/10000:,.1f} 万", f"电商+B2B+门店 三渠道合计"),
    ("Top3 品类集中度", f"{cat['top3_share_pct']:.2f}%", "定西宽粉 / 生鲜水果 / 高原夏菜"),
    ("电商复购率", f"{usr['repurchase_rate_pct']:.2f}%", f"{usr['buyers']:,} 购买用户 · 人均 {usr['avg_freq']:.2f} 次"),
    ("最优品类 ROI", f"{cat['best_roi'][0]['ROI']:.2f}", f"{cat['best_roi'][0]['品类']}（毛利/内容投入）"),
    ("新模式履约时效", f"{float(new_mode['平均履约时效_天']):.2f} 天", f"旧模式 {float(old_mode['平均履约时效_天']):.2f} 天 · 缩短 {float(old_mode['平均履约时效_天'])-float(new_mode['平均履约时效_天']):.1f} 天"),
    ("新模式配送损耗率", f"{float(new_mode['配送损耗率_pct']):.2f}%", f"旧模式 {float(old_mode['配送损耗率_pct']):.2f}% · 下降 {float(old_mode['配送损耗率_pct'])-float(new_mode['配送损耗率_pct']):.2f}pct"),
    ("门店合计月净利", f"¥{sto['合计月净利']:,.0f}", f"4 家门店均未达盈亏平衡 · 缺口 ¥{sto['月销售缺口合计']:,.0f}/月"),
    ("客流+60% 情景净利", f"¥{sto['情景合计净利_客流加60pct']:,.0f}/月", "情景测算：客流提升 60% 后整体转正"),
]
kpi_start = row
for idx, (label, value, note) in enumerate(kpis):
    r = kpi_start + (idx // 4) * 4
    col0 = 1 + (idx % 4) * 2
    ws.merge_cells(start_row=r, start_column=col0, end_row=r, end_column=col0 + 1)
    ws.merge_cells(start_row=r + 1, start_column=col0, end_row=r + 1, end_column=col0 + 1)
    ws.merge_cells(start_row=r + 2, start_column=col0, end_row=r + 2, end_column=col0 + 1)
    cl = ws.cell(row=r, column=col0, value=label)
    cl.font = F_KPI_LABEL
    cl.alignment = CENTER
    cv = ws.cell(row=r + 1, column=col0, value=value)
    cv.font = F_KPI_NUM if "净利" not in label or "+" in label else F_WARN
    cv.alignment = CENTER
    cn = ws.cell(row=r + 2, column=col0, value=note)
    cn.font = Font(name="微软雅黑", size=8, color="888888")
    cn.alignment = CENTER
    for rr in range(r, r + 3):
        for cc in range(col0, col0 + 2):
            ws.cell(row=rr, column=cc).fill = C_KPI_BG
            ws.cell(row=rr, column=cc).border = BORDER
    ws.row_dimensions[r].height = 16
    ws.row_dimensions[r + 1].height = 28
    ws.row_dimensions[r + 2].height = 24

row = kpi_start + 8 + 1
row = write_section_title(ws, row, "🧭 分析结论速览")
conclusions = [
    "1. 品类：定西宽粉及薯制品为绝对支柱（GMV 占比 26.44%、毛利率 53.93%），礼盒组合 ROI 最高（1.68）但体量小，建议作为节礼场景重点培育。",
    "2. 用户：复购率 22.14%，重要价值客户以 14.11% 的人数贡献 26.87% 的电商金额；近半数用户为低频挽留型，会员召回空间大。",
    "3. 直播：漏斗最大流失在「曝光→进入」（进入率仅 9.9%）与「互动→下单」（5.1%）；傍晚场 UV 价值 1.01 元最高，黄金场贡献 GMV 最多。",
    "4. 供应链：集中仓储+分级配送新模式使电商订单时效 3.76→2.46 天、损耗 1.51%→0.87%、单位运费 3.16→1.79 元/kg，全链路改善显著。",
    "5. 门店：县域成本结构下 4 家门店月销售均低于盈亏平衡线（合计缺口 3.33 万/月），客流需提升约 60% 方可整体打平，建议先做引流验证再扩张。",
]
for t in conclusions:
    row = write_note(ws, row, t, Font(name="微软雅黑", size=10, color="333333"))
row += 1
row = write_note(ws, row, f"数据生成时间：{R['generated_at']} ｜ SQL ↔ Python 复算交叉校验：{'通过 ✔' if R['cross_check_passed'] else '未通过 ✘'}")
sheet_footer(ws, row)
for col in range(1, 11):
    ws.column_dimensions[get_column_letter(col)].width = 13

# ================================================================ Sheet2 品类ROI
ws = wb.create_sheet("品类ROI")
ws.sheet_view.showGridLines = False
row = 1
row = write_section_title(ws, row, "① 品类结构与 ROI 分析（电商 + B2B + 门店 全渠道口径）")
row = write_note(ws, row, "品类ROI = 品类毛利 ÷ 内容投入成本；内容投入含直播/短视频/图文等带货投入。对应图1、图2。")
df_show = df_cat.copy()
row = write_df(ws, df_show, row, roi_col="品类ROI") + 1
top3_txt = "、".join(f"{t['品类']}（{t['占比_pct']}%）" for t in cat["top3"])
row = write_note(ws, row, f"▸ Top3 品类 {top3_txt}，合计占全渠道 GMV 的 {cat['top3_share_pct']}%。", F_BOLD)
row = write_note(ws, row, f"▸ ROI≥1 的品类：{'、'.join(r['品类'] for _, r in df_cat[df_cat['品类ROI'] >= 1].iterrows())}；ROI<1 品类的内容投放需复盘。")
row += 1
used = add_image(ws, "fig1_品类GMV帕累托.png", row)
row += used
used = add_image(ws, "fig2_品类ROI矩阵.png", row)
row += used
sheet_footer(ws, row)
ws.freeze_panes = "A4"

# ================================================================ Sheet3 用户分层
ws = wb.create_sheet("用户分层")
ws.sheet_view.showGridLines = False
row = 1
row = write_section_title(ws, row, "② 用户 RFM 分层与复购分析（电商渠道）")
row = write_note(ws, row, "分层口径：R 按最近购买天数中位数二分，F 按购买频次≥2 二分，M 按累计金额中位数二分（NTILE）。对应图3。")
row = write_note(ws, row, f"▸ 购买用户 {usr['buyers']:,} 人，复购率 {usr['repurchase_rate_pct']:.2f}%，人均购买 {usr['avg_freq']:.2f} 次。", F_BOLD)
df_show = df_rfm.copy()
for c_ in ["人群占比_pct", "平均R天数", "平均频次", "平均金额", "累计贡献金额", "贡献占比_pct"]:
    if c_ in df_show.columns:
        df_show[c_] = pd.to_numeric(df_show[c_], errors="coerce")
row = write_df(ws, df_show, row) + 1
seg = sorted(usr["segments"], key=lambda x: -float(x["贡献占比_pct"]))[0]
row = write_note(ws, row, f"▸ 「{seg['人群分层']}」人数占比 {seg['人群占比_pct']}%，贡献电商金额 {seg['贡献占比_pct']}%（¥{float(seg['累计贡献金额']):,.0f}），是运营基本盘。")
row = write_note(ws, row, "▸ 挽留/发展类客户合计人数占比过半但贡献不足 25%，建议优惠券召回 + 组合装提客单。")
row += 1
used = add_image(ws, "fig3_RFM用户分层.png", row)
row += used
sheet_footer(ws, row)
ws.freeze_panes = "A5"

# ================================================================ Sheet4 直播漏斗
ws = wb.create_sheet("直播漏斗")
ws.sheet_view.showGridLines = False
row = 1
row = write_section_title(ws, row, "③ 直播转化漏斗与时段效果")
row = write_note(ws, row, "漏斗为六段口径：曝光→进入→停留→互动→下单→付款（基于直播分钟级流量明细汇总）。对应图4、图5。")
df_show = df_funnel.copy()
df_show["人次"] = pd.to_numeric(df_show["人次"], errors="coerce")
row = write_df(ws, df_show, row) + 1
steps = df_show["人次"].tolist()
row = write_note(ws, row, f"▸ 进入率 {steps[1]/steps[0]*100:.1f}%、互动→下单转化 {steps[4]/steps[3]*100:.1f}%、下单→付款 {steps[5]/steps[4]*100:.1f}%：流量引入与货品讲解转化是两大优化点。", F_BOLD)
row += 1
row = write_section_title(ws, row, "分时段效果")
df_show = df_slot.copy()
for c_ in ["场均曝光", "场均场观", "场均付款人数", "累计GMV", "UV价值", "平均进入率"]:
    if c_ in df_show.columns:
        df_show[c_] = pd.to_numeric(df_show[c_], errors="coerce")
row = write_df(ws, df_show, row) + 1
best = max(R["funnel"]["slots"], key=lambda x: float(x["UV价值"]))
row = write_note(ws, row, f"▸ 「{best['时段']}」UV 价值最高（{float(best['UV价值']):.2f} 元），黄金场（19:00-20:00）累计 GMV 最大，建议稳黄金场、加傍晚场。")
row += 1
used = add_image(ws, "fig4_直播漏斗.png", row)
row += used
used = add_image(ws, "fig5_直播时段效果.png", row)
row += used
sheet_footer(ws, row)
ws.freeze_panes = "A4"

# ================================================================ Sheet5 供应链
ws = wb.create_sheet("供应链")
ws.sheet_view.showGridLines = False
row = 1
row = write_section_title(ws, row, "④ 供应链：新旧配送模式对比与分环节损耗")
row = write_note(ws, row, "新模式 = 集中仓储 + 分级配送；旧模式 = 分散直发 + 统配。对应图6、图7。")
df_show = df_mode.copy()
for c_ in ["平均履约时效_天", "配送损耗率_pct", "单位运费_元每kg", "累计运费"]:
    df_show[c_] = pd.to_numeric(df_show[c_], errors="coerce")
row = write_df(ws, df_show, row) + 1
row = write_note(ws, row, f"▸ 电商订单：时效 {float(old_mode['平均履约时效_天']):.2f}→{float(new_mode['平均履约时效_天']):.2f} 天，"
                          f"损耗 {float(old_mode['配送损耗率_pct']):.2f}%→{float(new_mode['配送损耗率_pct']):.2f}%，"
                          f"单位运费 {float(old_mode['单位运费_元每kg']):.2f}→{float(new_mode['单位运费_元每kg']):.2f} 元/kg。", F_BOLD)
row += 1
row = write_section_title(ws, row, "分环节损耗率")
df_show = df_loss.copy()
df_show["损耗率_pct"] = pd.to_numeric(df_show["损耗率_pct"], errors="coerce")
df_show["损耗量_kg"] = pd.to_numeric(df_show["损耗量_kg"], errors="coerce")
row = write_df(ws, df_show, row) + 1
row += 1
used = add_image(ws, "fig6_配送模式对比.png", row)
row += used
used = add_image(ws, "fig7_分环节损耗.png", row)
row += used
sheet_footer(ws, row)
ws.freeze_panes = "A4"

# ================================================================ Sheet6 门店测算
ws = wb.create_sheet("门店测算")
ws.sheet_view.showGridLines = False
row = 1
row = write_section_title(ws, row, "⑤ 门店盈利测算与盈亏平衡分析（县域成本口径）")
row = write_note(ws, row, "月净利 = 月均毛利 −（租金 + 人工 + 水电 + 折旧摊销）；盈亏平衡月销售额 = 月固定成本 ÷ 综合毛利率。对应图8、图9。")
df_show = df_store.copy()
row = write_df(ws, df_show, row, warn_neg_col="月净利") + 1
row = write_note(ws, row, f"▸ 4 家门店现状均未达盈亏平衡：合计月销售 ¥{sto['现状月销售合计']:,.0f} vs 盈亏平衡要求 ¥{sto['盈亏平衡月销售合计']:,.0f}，缺口 ¥{sto['月销售缺口合计']:,.0f}/月。", F_WARN)
row = write_note(ws, row, f"▸ 情景测算：客流 +30% 合计净利 ¥{sto['情景合计净利_客流加30pct']:,.0f}/月；+60% ¥{sto['情景合计净利_客流加60pct']:,.0f}/月（转正）；翻倍 ¥{sto['情景合计净利_客流翻倍']:,.0f}/月。", F_GOOD)
row += 1
row = write_section_title(ws, row, "门店汇总")
row = write_df(ws, df_store_sum, row) + 2
used = add_image(ws, "fig8_门店盈亏平衡.png", row)
row += used
used = add_image(ws, "fig9_门店情景敏感性.png", row)
row += used
sheet_footer(ws, row)
ws.freeze_panes = "A4"

# ---------------------------------------------------------------- 保存
wb.save(OUT_PATH)
print(f"✔ Excel 看板已生成：{OUT_PATH}")
print(f"  共 6 个 sheet：{', '.join(wb.sheetnames)}")
print(f"  （模拟数据，仅用于方法复现，不代表真实经营结果）")
